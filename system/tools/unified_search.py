#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# SPDX-License-Identifier: MIT
"""
Tool: unified_search
Version: 1.0.0
Author: Claude Opus 4.6
Created: 2026-02-20
Implements: SQ064 (Semantische Suche) + SQ047 (Wissensindexierung)

Description:
    Unified search engine for BACH. Combines FTS5 full-text search across all
    BACH data sources: documents, wiki, memory (working/facts/lessons).
    Includes ProFiler-inspired file discovery with hash-based dedup and
    tags derived from path structure.

    Future: sqlite-vec vector search (cosine + BM25 hybrid).

Architecture:
    search_index (table)  <-- unified metadata + content from all sources
    search_fts   (FTS5)   <-- full-text index synced via triggers
    search_tags  (table)  <-- tags per indexed item (ProFiler pattern)
"""

__version__ = "1.0.0"
__author__ = "Claude Opus 4.6"

import sqlite3
import hashlib
import os
import sys
import json
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Optional, Tuple, Any

# BACH Root
BACH_ROOT = Path(__file__).parent.parent
DB_PATH = BACH_ROOT / "data" / "bach.db"

# --- File type mappings (from DocumentIndexer + ProFiler) ---

SUPPORTED_TEXT_EXTENSIONS = {
    '.txt', '.md', '.py', '.js', '.ts', '.json', '.csv', '.xml',
    '.html', '.htm', '.log', '.cfg', '.ini', '.yaml', '.yml',
    '.rst', '.toml', '.sh', '.bat', '.ps1', '.sql',
}

FILE_CATEGORIES = {
    '.txt': 'dokument', '.md': 'dokument', '.pdf': 'dokument',
    '.docx': 'dokument', '.rst': 'dokument',
    '.py': 'code', '.js': 'code', '.ts': 'code', '.java': 'code',
    '.cpp': 'code', '.c': 'code', '.h': 'code', '.sh': 'code',
    '.sql': 'code', '.bat': 'code', '.ps1': 'code',
    '.json': 'daten', '.csv': 'daten', '.xml': 'daten',
    '.yaml': 'daten', '.yml': 'daten', '.toml': 'daten',
    '.jpg': 'bild', '.jpeg': 'bild', '.png': 'bild', '.gif': 'bild',
    '.mp3': 'audio', '.wav': 'audio', '.mp4': 'video', '.avi': 'video',
    '.zip': 'archiv', '.tar': 'archiv', '.gz': 'archiv',
}

# --- Schema ---

SCHEMA_SQL = """
-- Unified search index: one row per indexed item from any source
CREATE TABLE IF NOT EXISTS search_index (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    source TEXT NOT NULL,          -- 'document', 'wiki', 'memory_working', 'memory_fact', 'memory_lesson', 'file'
    source_id TEXT NOT NULL,       -- ID/key in source table
    source_path TEXT,              -- file path or DB reference
    title TEXT NOT NULL,           -- display title
    content TEXT,                  -- full text content (truncated to 100k)
    category TEXT,                 -- file category or content type
    content_hash TEXT,             -- SHA256 for dedup
    word_count INTEGER DEFAULT 0,
    file_size INTEGER DEFAULT 0,
    indexed_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    UNIQUE(source, source_id)
);

-- Tags table (ProFiler pattern: tags derived from path structure)
CREATE TABLE IF NOT EXISTS search_tags (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    search_id INTEGER NOT NULL REFERENCES search_index(id) ON DELETE CASCADE,
    tag TEXT NOT NULL,
    UNIQUE(search_id, tag)
);

-- FTS5 virtual table over search_index
CREATE VIRTUAL TABLE IF NOT EXISTS search_fts USING fts5(
    title,
    content,
    tokenize='unicode61'
);

-- Auto-sync triggers: search_index <-> search_fts
CREATE TRIGGER IF NOT EXISTS search_idx_ai AFTER INSERT ON search_index BEGIN
    INSERT INTO search_fts(rowid, title, content)
    VALUES (new.id, new.title, COALESCE(new.content, ''));
END;

CREATE TRIGGER IF NOT EXISTS search_idx_ad AFTER DELETE ON search_index BEGIN
    INSERT INTO search_fts(search_fts, rowid, title, content)
    VALUES ('delete', old.id, old.title, COALESCE(old.content, ''));
END;

CREATE TRIGGER IF NOT EXISTS search_idx_au AFTER UPDATE ON search_index BEGIN
    INSERT INTO search_fts(search_fts, rowid, title, content)
    VALUES ('delete', old.id, old.title, COALESCE(old.content, ''));
    INSERT INTO search_fts(rowid, title, content)
    VALUES (new.id, new.title, COALESCE(new.content, ''));
END;

-- Index for fast tag lookups
CREATE INDEX IF NOT EXISTS idx_search_tags_tag ON search_tags(tag);
CREATE INDEX IF NOT EXISTS idx_search_index_source ON search_index(source);
CREATE INDEX IF NOT EXISTS idx_search_index_hash ON search_index(content_hash);
"""


def _is_cloud_placeholder(path: str) -> bool:
    """Detect OneDrive cloud placeholders (ProFiler pattern)."""
    if os.name != 'nt':
        return False
    try:
        attrs = os.stat(path).st_file_attributes
        return bool((attrs & 0x1000) or (attrs & 0x400))
    except (OSError, AttributeError):
        return False


def _sha256_file(path: str, chunk_size: int = 1024 * 1024) -> Optional[str]:
    """SHA256 hash of a file (ProFiler pattern)."""
    h = hashlib.sha256()
    try:
        with open(path, "rb") as f:
            while True:
                chunk = f.read(chunk_size)
                if not chunk:
                    break
                h.update(chunk)
    except (PermissionError, OSError):
        return None
    return h.hexdigest()


def _sha256_text(text: str) -> str:
    """SHA256 hash of text content."""
    return hashlib.sha256(text.encode('utf-8', errors='ignore')).hexdigest()


def _path_to_tags(filepath: str, root: str) -> List[str]:
    """Extract tags from path structure (ProFiler pattern).

    Example: /docs/health/reports/blood.txt -> ['docs', 'health', 'reports']
    """
    try:
        rel = os.path.relpath(filepath, root)
        parts = os.path.dirname(rel).replace('\\', '/').split('/')
        return [p for p in parts if p and p != '.']
    except (ValueError, TypeError):
        return []


def _extract_text(file_path: Path) -> str:
    """Extract text content from a file. Supports text, PDF, DOCX."""
    ext = file_path.suffix.lower()

    if ext in SUPPORTED_TEXT_EXTENSIONS:
        try:
            return file_path.read_text(encoding='utf-8', errors='ignore')
        except Exception:
            try:
                return file_path.read_text(encoding='latin-1', errors='ignore')
            except Exception:
                return ""

    if ext == '.pdf':
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(file_path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except (ImportError, Exception):
            pass
        try:
            import pdfplumber
            with pdfplumber.open(str(file_path)) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        except (ImportError, Exception):
            pass
        return ""

    if ext == '.docx':
        try:
            from docx import Document
            doc = Document(str(file_path))
            return "\n".join(p.text for p in doc.paragraphs)
        except (ImportError, Exception):
            return ""

    return ""


class UnifiedSearch:
    """Unified search engine for BACH.

    Indexes and searches across:
    - document_index (existing FTS5, help/ docs)
    - wiki_articles (413 articles)
    - memory_working, memory_facts, memory_lessons
    - arbitrary directories (ProFiler-style file scan)
    """

    def __init__(self, db_path: Path = DB_PATH):
        self.db_path = db_path

    def _get_db(self) -> sqlite3.Connection:
        conn = sqlite3.connect(str(self.db_path))
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA foreign_keys=ON")
        return conn

    def ensure_schema(self):
        """Create search_index, search_fts, search_tags if not exist."""
        conn = self._get_db()
        try:
            conn.executescript(SCHEMA_SQL)
            conn.commit()
        finally:
            conn.close()

    # ------------------------------------------------------------------
    # INDEXING: Pull data from all BACH sources into search_index
    # ------------------------------------------------------------------

    def index_all(self) -> Tuple[bool, str]:
        """Rebuild entire unified index from all sources."""
        self.ensure_schema()
        results = []

        ok, msg = self.index_wiki()
        results.append(msg)

        ok, msg = self.index_memory()
        results.append(msg)

        ok, msg = self.index_documents()
        results.append(msg)

        ok, msg = self.index_knowledgedigest()
        results.append(msg)

        return True, "\n".join(results)

    def index_wiki(self) -> Tuple[bool, str]:
        """Index all wiki_articles into search_index."""
        conn = self._get_db()
        count = 0
        try:
            rows = conn.execute(
                "SELECT path, title, content, category, tags FROM wiki_articles"
            ).fetchall()

            for row in rows:
                content = row['content'] or ''
                title = row['title'] or Path(row['path']).stem
                content_hash = _sha256_text(content)
                word_count = len(content.split()) if content else 0
                tags_str = row['tags'] or ''

                conn.execute("""
                    INSERT INTO search_index
                        (source, source_id, source_path, title, content,
                         category, content_hash, word_count)
                    VALUES ('wiki', ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(source, source_id) DO UPDATE SET
                        title=excluded.title, content=excluded.content,
                        content_hash=excluded.content_hash,
                        word_count=excluded.word_count,
                        indexed_at=CURRENT_TIMESTAMP
                """, (
                    row['path'], row['path'], title,
                    content[:100000], row['category'] or 'wiki',
                    content_hash, word_count
                ))

                # Tags from wiki
                if tags_str:
                    search_id = conn.execute(
                        "SELECT id FROM search_index WHERE source='wiki' AND source_id=?",
                        (row['path'],)
                    ).fetchone()
                    if search_id:
                        for tag in tags_str.split(','):
                            tag = tag.strip()
                            if tag:
                                conn.execute(
                                    "INSERT OR IGNORE INTO search_tags (search_id, tag) VALUES (?, ?)",
                                    (search_id['id'], tag)
                                )
                count += 1

            conn.commit()
            return True, f"[Wiki] {count} Artikel indexiert"
        except Exception as e:
            return False, f"[Wiki] Fehler: {e}"
        finally:
            conn.close()

    def index_memory(self) -> Tuple[bool, str]:
        """Index memory_working, memory_facts, memory_lessons."""
        conn = self._get_db()
        counts = {'working': 0, 'facts': 0, 'lessons': 0}
        try:
            # memory_working
            rows = conn.execute(
                "SELECT id, type, content, tags FROM memory_working WHERE is_active=1"
            ).fetchall()
            for row in rows:
                content = row['content'] or ''
                title = f"[{row['type']}] {content[:60]}"
                conn.execute("""
                    INSERT INTO search_index
                        (source, source_id, title, content, category, content_hash, word_count)
                    VALUES ('memory_working', ?, ?, ?, 'memory', ?, ?)
                    ON CONFLICT(source, source_id) DO UPDATE SET
                        title=excluded.title, content=excluded.content,
                        content_hash=excluded.content_hash,
                        indexed_at=CURRENT_TIMESTAMP
                """, (
                    str(row['id']), title, content,
                    _sha256_text(content), len(content.split())
                ))
                counts['working'] += 1

            # memory_facts
            rows = conn.execute(
                "SELECT id, category, key, value, source FROM memory_facts"
            ).fetchall()
            for row in rows:
                value = row['value'] or ''
                title = f"{row['category']}/{row['key']}"
                content = f"{row['key']}: {value}"
                conn.execute("""
                    INSERT INTO search_index
                        (source, source_id, title, content, category, content_hash, word_count)
                    VALUES ('memory_fact', ?, ?, ?, 'memory', ?, ?)
                    ON CONFLICT(source, source_id) DO UPDATE SET
                        title=excluded.title, content=excluded.content,
                        content_hash=excluded.content_hash,
                        indexed_at=CURRENT_TIMESTAMP
                """, (
                    str(row['id']), title, content,
                    _sha256_text(content), len(content.split())
                ))
                counts['facts'] += 1

            # memory_lessons
            rows = conn.execute(
                "SELECT id, category, title, problem, solution, trigger_words FROM memory_lessons WHERE is_active=1"
            ).fetchall()
            for row in rows:
                parts = [row['title'] or '', row['problem'] or '', row['solution'] or '']
                content = "\n".join(p for p in parts if p)
                title = row['title'] or f"Lesson #{row['id']}"
                conn.execute("""
                    INSERT INTO search_index
                        (source, source_id, title, content, category, content_hash, word_count)
                    VALUES ('memory_lesson', ?, ?, ?, 'memory', ?, ?)
                    ON CONFLICT(source, source_id) DO UPDATE SET
                        title=excluded.title, content=excluded.content,
                        content_hash=excluded.content_hash,
                        indexed_at=CURRENT_TIMESTAMP
                """, (
                    str(row['id']), title, content,
                    _sha256_text(content), len(content.split())
                ))

                # Tags from trigger_words
                if row['trigger_words']:
                    search_id = conn.execute(
                        "SELECT id FROM search_index WHERE source='memory_lesson' AND source_id=?",
                        (str(row['id']),)
                    ).fetchone()
                    if search_id:
                        for tw in row['trigger_words'].split(','):
                            tw = tw.strip()
                            if tw:
                                conn.execute(
                                    "INSERT OR IGNORE INTO search_tags (search_id, tag) VALUES (?, ?)",
                                    (search_id['id'], tw)
                                )
                counts['lessons'] += 1

            conn.commit()
            return True, (
                f"[Memory] working={counts['working']}, "
                f"facts={counts['facts']}, lessons={counts['lessons']}"
            )
        except Exception as e:
            return False, f"[Memory] Fehler: {e}"
        finally:
            conn.close()

    def index_documents(self) -> Tuple[bool, str]:
        """Import existing document_index entries into unified search."""
        conn = self._get_db()
        count = 0
        try:
            rows = conn.execute("""
                SELECT id, file_path, file_name, file_category,
                       content_text, content_summary, file_hash, word_count, file_size
                FROM document_index
            """).fetchall()

            for row in rows:
                content = row['content_text'] or row['content_summary'] or ''
                title = row['file_name'] or 'Unbekannt'
                conn.execute("""
                    INSERT INTO search_index
                        (source, source_id, source_path, title, content,
                         category, content_hash, word_count, file_size)
                    VALUES ('document', ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(source, source_id) DO UPDATE SET
                        title=excluded.title, content=excluded.content,
                        content_hash=excluded.content_hash,
                        word_count=excluded.word_count,
                        indexed_at=CURRENT_TIMESTAMP
                """, (
                    str(row['id']), row['file_path'], title,
                    content[:100000], row['file_category'],
                    row['file_hash'], row['word_count'] or 0,
                    row['file_size'] or 0
                ))
                count += 1

            conn.commit()
            return True, f"[Dokumente] {count} aus document_index importiert"
        except Exception as e:
            return False, f"[Dokumente] Fehler: {e}"
        finally:
            conn.close()

    def index_knowledgedigest(self, kd_db_path=None) -> Tuple[bool, str]:
        """Index Skills und Wiki aus KnowledgeDigest knowledge.db.

        Indexiert alle Skills und Wiki-Artikel aus der KnowledgeDigest-Datenbank
        in den unified search index. Chunks werden pro Dokument konkateniert.

        Args:
            kd_db_path: Pfad zur knowledge.db (Default: MODULAR_AGENTS/KnowledgeDigest/data/knowledge.db)

        Returns:
            Tuple[bool, str]: (Erfolg, Statusmeldung)
        """
        # Lazy-Import: KnowledgeDigest aus MODULAR_AGENTS
        modular_agents_dir = BACH_ROOT.parent.parent / "MODULAR_AGENTS"
        kd_package_dir = modular_agents_dir / "KnowledgeDigest"

        if not kd_package_dir.exists():
            return True, "[KnowledgeDigest] Nicht gefunden, uebersprungen."

        # sys.path erweitern fuer Import
        kd_parent = str(modular_agents_dir)
        path_added = False
        if kd_parent not in sys.path:
            sys.path.insert(0, kd_parent)
            path_added = True

        try:
            from KnowledgeDigest import KnowledgeDigest
        except ImportError as e:
            if path_added:
                sys.path.remove(kd_parent)
            return True, f"[KnowledgeDigest] Import fehlgeschlagen ({e}), uebersprungen."

        # knowledge.db Pfad
        if kd_db_path is None:
            kd_db_path = kd_package_dir / "data" / "knowledge.db"
        kd_db_path = Path(kd_db_path)
        if not kd_db_path.exists():
            return True, f"[KnowledgeDigest] knowledge.db nicht gefunden: {kd_db_path}"

        conn = self._get_db()
        skill_count = 0
        wiki_count = 0

        try:
            kd = KnowledgeDigest(db_path=kd_db_path)

            # Alte KnowledgeDigest-Eintraege loeschen
            # FTS5 AFTER DELETE Trigger muss temporaer deaktiviert werden,
            # da der 'delete'-Befehl bei grossen Texten fehlschlagen kann.
            existing_count = conn.execute(
                "SELECT COUNT(*) FROM search_index "
                "WHERE source IN ('knowledgedigest_skill', 'knowledgedigest_wiki')"
            ).fetchone()[0]

            if existing_count > 0:
                conn.execute("DROP TRIGGER IF EXISTS search_idx_ad")
                conn.execute(
                    "DELETE FROM search_tags WHERE search_id IN "
                    "(SELECT id FROM search_index WHERE source IN "
                    "('knowledgedigest_skill', 'knowledgedigest_wiki'))"
                )
                conn.execute(
                    "DELETE FROM search_index "
                    "WHERE source IN ('knowledgedigest_skill', 'knowledgedigest_wiki')"
                )
                conn.execute("INSERT INTO search_fts(search_fts) VALUES('rebuild')")
                conn.executescript("""
                    CREATE TRIGGER IF NOT EXISTS search_idx_ad AFTER DELETE ON search_index BEGIN
                        INSERT INTO search_fts(search_fts, rowid, title, content)
                        VALUES ('delete', old.id, old.title, COALESCE(old.content, ''));
                    END;
                """)
                conn.commit()

            # --- Skills indexieren ---
            skills = kd.list_skills(limit=2000)
            for skill_info in skills:
                skill_name = skill_info['skill_name']
                skill = kd.get_skill(skill_name)
                if not skill:
                    continue

                # Chunks konkatenieren
                chunks = skill.get('chunks', [])
                content = "\n\n".join(c['content'] for c in chunks if c.get('content'))
                if not content:
                    continue

                title = skill_name
                source_id = f"skill:{skill_name}"
                source_path = skill.get('path', '')
                category = skill.get('category', 'skill')
                content_hash = _sha256_text(content)
                word_count = len(content.split())
                description = skill.get('description', '')
                if description:
                    content = f"{description}\n\n{content}"

                conn.execute("""
                    INSERT INTO search_index
                        (source, source_id, source_path, title, content,
                         category, content_hash, word_count)
                    VALUES ('knowledgedigest_skill', ?, ?, ?, ?, ?, ?, ?)
                """, (
                    source_id, source_path, title,
                    content[:100000], category,
                    content_hash, word_count
                ))

                # Tags aus Keywords
                keywords = skill.get('keywords', [])
                if keywords:
                    search_id = conn.execute(
                        "SELECT id FROM search_index WHERE source='knowledgedigest_skill' AND source_id=?",
                        (source_id,)
                    ).fetchone()
                    if search_id:
                        for kw in keywords:
                            kw = str(kw).strip()
                            if kw and len(kw) <= 50 and not kw.startswith('-' * 3):
                                conn.execute(
                                    "INSERT OR IGNORE INTO search_tags (search_id, tag) VALUES (?, ?)",
                                    (search_id['id'], kw)
                                )

                skill_count += 1
                if skill_count % 100 == 0:
                    conn.commit()

            conn.commit()

            # --- Wiki indexieren ---
            wikis = kd.list_wikis(limit=2000)
            for wiki_info in wikis:
                wiki_path = wiki_info['wiki_path']
                wiki = kd.get_wiki(wiki_path)
                if not wiki:
                    continue

                # Chunks konkatenieren
                chunks = wiki.get('chunks', [])
                content = "\n\n".join(c['content'] for c in chunks if c.get('content'))
                if not content:
                    continue

                title = wiki.get('title', '') or Path(wiki_path).stem
                source_id = f"wiki:{wiki_path}"
                category = wiki.get('category', 'wiki')
                content_hash = _sha256_text(content)
                word_count = len(content.split())

                conn.execute("""
                    INSERT INTO search_index
                        (source, source_id, source_path, title, content,
                         category, content_hash, word_count)
                    VALUES ('knowledgedigest_wiki', ?, ?, ?, ?, ?, ?, ?)
                """, (
                    source_id, wiki_path, title,
                    content[:100000], category,
                    content_hash, word_count
                ))

                # Tags aus Keywords
                keywords = wiki.get('keywords', [])
                if keywords:
                    search_id = conn.execute(
                        "SELECT id FROM search_index WHERE source='knowledgedigest_wiki' AND source_id=?",
                        (source_id,)
                    ).fetchone()
                    if search_id:
                        for kw in keywords:
                            kw = str(kw).strip()
                            if kw:
                                conn.execute(
                                    "INSERT OR IGNORE INTO search_tags (search_id, tag) VALUES (?, ?)",
                                    (search_id['id'], kw)
                                )

                wiki_count += 1
                if wiki_count % 100 == 0:
                    conn.commit()

            conn.commit()

            try:
                kd.close()
            except Exception:
                pass

            return True, f"[KnowledgeDigest] {skill_count} Skills + {wiki_count} Wiki indexiert"

        except Exception as e:
            return False, f"[KnowledgeDigest] Fehler: {e}"
        finally:
            conn.close()
            if path_added and kd_parent in sys.path:
                sys.path.remove(kd_parent)

    def scan_directory(self, directory: str, tags_from_path: bool = True,
                       recursive: bool = True) -> Tuple[bool, str]:
        """ProFiler-style directory scan: index files with hash dedup and path tags.

        This extends the index with files from arbitrary directories,
        not just BACH's own docs/wiki.
        """
        self.ensure_schema()
        dir_path = Path(directory)
        if not dir_path.exists():
            return False, f"Ordner nicht gefunden: {directory}"

        conn = self._get_db()
        count = 0
        skipped = 0
        errors = 0

        try:
            pattern = dir_path.rglob("*") if recursive else dir_path.glob("*")
            for filepath in pattern:
                if not filepath.is_file():
                    continue
                if any(part.startswith('.') for part in filepath.parts):
                    continue

                ext = filepath.suffix.lower()
                category = FILE_CATEGORIES.get(ext, 'sonstiges')

                try:
                    stat = filepath.stat()

                    # Skip cloud placeholders
                    if _is_cloud_placeholder(str(filepath)):
                        skipped += 1
                        continue

                    # Hash for dedup
                    content_hash = _sha256_file(str(filepath))
                    if not content_hash:
                        skipped += 1
                        continue

                    # Check if already indexed with same hash
                    existing = conn.execute(
                        "SELECT id, content_hash FROM search_index "
                        "WHERE source='file' AND source_path=?",
                        (str(filepath),)
                    ).fetchone()

                    if existing and existing['content_hash'] == content_hash:
                        skipped += 1
                        continue

                    # Extract text for searchable types
                    content = ""
                    if ext in SUPPORTED_TEXT_EXTENSIONS or ext in ('.pdf', '.docx'):
                        content = _extract_text(filepath)

                    word_count = len(content.split()) if content else 0
                    title = filepath.name

                    conn.execute("""
                        INSERT INTO search_index
                            (source, source_id, source_path, title, content,
                             category, content_hash, word_count, file_size)
                        VALUES ('file', ?, ?, ?, ?, ?, ?, ?, ?)
                        ON CONFLICT(source, source_id) DO UPDATE SET
                            title=excluded.title, content=excluded.content,
                            content_hash=excluded.content_hash,
                            word_count=excluded.word_count,
                            file_size=excluded.file_size,
                            indexed_at=CURRENT_TIMESTAMP
                    """, (
                        str(filepath), str(filepath), title,
                        content[:100000] if content else None,
                        category, content_hash, word_count, stat.st_size
                    ))

                    # Tags from path structure
                    if tags_from_path:
                        search_id = conn.execute(
                            "SELECT id FROM search_index WHERE source='file' AND source_id=?",
                            (str(filepath),)
                        ).fetchone()
                        if search_id:
                            for tag in _path_to_tags(str(filepath), directory):
                                conn.execute(
                                    "INSERT OR IGNORE INTO search_tags (search_id, tag) VALUES (?, ?)",
                                    (search_id['id'], tag)
                                )

                    count += 1
                    if count % 100 == 0:
                        conn.commit()

                except (PermissionError, OSError):
                    errors += 1
                    continue

            conn.commit()
            return True, (
                f"[Scan] {directory}\n"
                f"  Indexiert: {count} | Uebersprungen: {skipped} | Fehler: {errors}"
            )
        except Exception as e:
            return False, f"[Scan] Fehler: {e}"
        finally:
            conn.close()

    # ------------------------------------------------------------------
    # SEARCH: Unified FTS5 search across all sources
    # ------------------------------------------------------------------

    def search(self, query: str, sources: Optional[List[str]] = None,
               tags: Optional[List[str]] = None, category: Optional[str] = None,
               limit: int = 20) -> Tuple[bool, List[Dict[str, Any]]]:
        """Search across all indexed sources using FTS5.

        Args:
            query: Search terms (FTS5 syntax supported)
            sources: Filter by source ('wiki', 'document', 'memory_*', 'file')
            tags: Filter by tags (AND logic)
            category: Filter by category
            limit: Max results

        Returns:
            (success, list of result dicts)
        """
        self.ensure_schema()
        conn = self._get_db()

        try:
            # Build query with filters
            params = [query, limit]
            where_clauses = []

            if sources:
                placeholders = ','.join('?' * len(sources))
                where_clauses.append(f"si.source IN ({placeholders})")
                params = [query] + sources + [limit]

            if category:
                where_clauses.append("si.category = ?")
                params.insert(-1, category)

            where_sql = ""
            if where_clauses:
                where_sql = "AND " + " AND ".join(where_clauses)

            # Tag filter via subquery
            tag_sql = ""
            if tags:
                tag_placeholders = ','.join('?' * len(tags))
                tag_sql = f"""
                    AND si.id IN (
                        SELECT search_id FROM search_tags
                        WHERE tag IN ({tag_placeholders})
                        GROUP BY search_id
                        HAVING COUNT(DISTINCT tag) = ?
                    )
                """
                # Insert tag params before limit
                limit_param = params.pop()
                params.extend(tags)
                params.append(len(tags))
                params.append(limit_param)

            sql = f"""
                SELECT
                    si.id, si.source, si.source_id, si.source_path,
                    si.title, si.category, si.word_count, si.file_size,
                    si.indexed_at,
                    snippet(search_fts, 1, '>>>', '<<<', '...', 30) as snippet,
                    search_fts.rank as relevance
                FROM search_fts
                JOIN search_index si ON search_fts.rowid = si.id
                WHERE search_fts MATCH ?
                {where_sql}
                {tag_sql}
                ORDER BY search_fts.rank
                LIMIT ?
            """

            rows = conn.execute(sql, params).fetchall()

            results = []
            for row in rows:
                # Fetch tags for this result
                tag_rows = conn.execute(
                    "SELECT tag FROM search_tags WHERE search_id=?",
                    (row['id'],)
                ).fetchall()

                results.append({
                    'id': row['id'],
                    'source': row['source'],
                    'source_id': row['source_id'],
                    'source_path': row['source_path'],
                    'title': row['title'],
                    'category': row['category'],
                    'word_count': row['word_count'],
                    'file_size': row['file_size'],
                    'snippet': row['snippet'],
                    'relevance': row['relevance'],
                    'tags': [r['tag'] for r in tag_rows],
                    'indexed_at': row['indexed_at'],
                })

            return True, results

        except Exception as e:
            # Fallback: LIKE search if FTS5 fails (empty index, bad query)
            try:
                like_param = f"%{query}%"
                rows = conn.execute("""
                    SELECT id, source, source_id, source_path, title,
                           category, word_count, file_size, indexed_at
                    FROM search_index
                    WHERE title LIKE ? OR content LIKE ?
                    LIMIT ?
                """, (like_param, like_param, limit)).fetchall()

                results = []
                for row in rows:
                    results.append({
                        'id': row['id'],
                        'source': row['source'],
                        'source_id': row['source_id'],
                        'source_path': row['source_path'],
                        'title': row['title'],
                        'category': row['category'],
                        'word_count': row['word_count'],
                        'file_size': row['file_size'],
                        'snippet': '(LIKE-Fallback)',
                        'relevance': 0,
                        'tags': [],
                        'indexed_at': row['indexed_at'],
                    })
                return True, results
            except Exception as e2:
                return False, []
        finally:
            conn.close()

    def search_by_tag(self, tags: List[str], limit: int = 20) -> Tuple[bool, List[Dict]]:
        """Find all items matching given tags (AND logic)."""
        self.ensure_schema()
        conn = self._get_db()
        try:
            placeholders = ','.join('?' * len(tags))
            rows = conn.execute(f"""
                SELECT si.*, GROUP_CONCAT(st.tag) as all_tags
                FROM search_index si
                JOIN search_tags st ON st.search_id = si.id
                WHERE st.tag IN ({placeholders})
                GROUP BY si.id
                HAVING COUNT(DISTINCT st.tag) = ?
                ORDER BY si.indexed_at DESC
                LIMIT ?
            """, tags + [len(tags), limit]).fetchall()

            results = [{
                'id': r['id'], 'source': r['source'],
                'source_id': r['source_id'], 'source_path': r['source_path'],
                'title': r['title'], 'category': r['category'],
                'tags': (r['all_tags'] or '').split(','),
            } for r in rows]

            return True, results
        except Exception as e:
            return False, []
        finally:
            conn.close()

    def list_tags(self, limit: int = 50) -> Tuple[bool, List[Dict]]:
        """List all tags with counts."""
        self.ensure_schema()
        conn = self._get_db()
        try:
            rows = conn.execute("""
                SELECT tag, COUNT(*) as count
                FROM search_tags
                GROUP BY tag
                ORDER BY count DESC
                LIMIT ?
            """, (limit,)).fetchall()
            return True, [{'tag': r['tag'], 'count': r['count']} for r in rows]
        except Exception as e:
            return False, []
        finally:
            conn.close()

    # ------------------------------------------------------------------
    # STATUS & MAINTENANCE
    # ------------------------------------------------------------------

    def status(self) -> Tuple[bool, Dict[str, Any]]:
        """Return index statistics."""
        self.ensure_schema()
        conn = self._get_db()
        try:
            total = conn.execute("SELECT COUNT(*) FROM search_index").fetchone()[0]

            by_source = conn.execute("""
                SELECT source, COUNT(*) as cnt, SUM(word_count) as words
                FROM search_index GROUP BY source ORDER BY cnt DESC
            """).fetchall()

            total_tags = conn.execute("SELECT COUNT(DISTINCT tag) FROM search_tags").fetchone()[0]
            total_words = conn.execute("SELECT SUM(word_count) FROM search_index").fetchone()[0] or 0

            # Check if search_fts is populated
            try:
                fts_count = conn.execute("SELECT COUNT(*) FROM search_fts").fetchone()[0]
            except Exception:
                fts_count = 0

            stats = {
                'total_items': total,
                'total_words': total_words,
                'total_tags': total_tags,
                'fts_entries': fts_count,
                'by_source': {r['source']: {'count': r['cnt'], 'words': r['words'] or 0} for r in by_source},
            }
            return True, stats
        except Exception as e:
            return False, {'error': str(e)}
        finally:
            conn.close()

    def clear(self) -> Tuple[bool, str]:
        """Clear the entire unified search index."""
        conn = self._get_db()
        try:
            conn.execute("DELETE FROM search_tags")
            conn.execute("DELETE FROM search_index")
            conn.commit()
            return True, "Unified Search Index geleert."
        except Exception as e:
            return False, f"Fehler beim Leeren: {e}"
        finally:
            conn.close()

    def rebuild(self) -> Tuple[bool, str]:
        """Clear and rebuild the entire index."""
        ok, msg = self.clear()
        if not ok:
            return ok, msg
        return self.index_all()

    def find_duplicates(self) -> Tuple[bool, List[Dict]]:
        """Find files with identical content (same hash, different paths)."""
        self.ensure_schema()
        conn = self._get_db()
        try:
            rows = conn.execute("""
                SELECT content_hash, COUNT(*) as cnt,
                       GROUP_CONCAT(source_path, ' | ') as paths
                FROM search_index
                WHERE content_hash IS NOT NULL
                  AND source IN ('file', 'document')
                GROUP BY content_hash
                HAVING cnt > 1
                ORDER BY cnt DESC
                LIMIT 50
            """).fetchall()

            dupes = [{
                'hash': r['content_hash'],
                'count': r['cnt'],
                'paths': r['paths'].split(' | ') if r['paths'] else [],
            } for r in rows]

            return True, dupes
        except Exception as e:
            return False, []
        finally:
            conn.close()


# --- CLI entrypoint ---

def main():
    import argparse
    parser = argparse.ArgumentParser(description="BACH Unified Search")
    sub = parser.add_subparsers(dest="cmd")

    sub.add_parser("index", help="Alle Quellen indexieren")
    sub.add_parser("rebuild", help="Index leeren und neu aufbauen")
    sub.add_parser("status", help="Index-Status")

    s = sub.add_parser("search", help="Suche")
    s.add_argument("query")
    s.add_argument("--source", "-s", help="Quelle filtern")
    s.add_argument("--tag", "-t", action="append", help="Tag-Filter")
    s.add_argument("--limit", "-l", type=int, default=20)

    sc = sub.add_parser("scan", help="Verzeichnis scannen")
    sc.add_argument("directory")
    sc.add_argument("--no-tags", action="store_true")

    sub.add_parser("tags", help="Alle Tags anzeigen")
    sub.add_parser("dupes", help="Duplikate finden")

    args = parser.parse_args()
    engine = UnifiedSearch()

    if args.cmd == "index":
        ok, msg = engine.index_all()
        print(msg)
    elif args.cmd == "rebuild":
        ok, msg = engine.rebuild()
        print(msg)
    elif args.cmd == "status":
        ok, stats = engine.status()
        print(json.dumps(stats, indent=2, ensure_ascii=False))
    elif args.cmd == "search":
        sources = [args.source] if args.source else None
        ok, results = engine.search(args.query, sources=sources, tags=args.tag, limit=args.limit)
        for r in results:
            tags = f" [{', '.join(r['tags'])}]" if r.get('tags') else ""
            print(f"  [{r['source']}] {r['title']}{tags}")
            if r.get('snippet'):
                print(f"    {r['snippet'][:120]}")
            print()
    elif args.cmd == "scan":
        ok, msg = engine.scan_directory(args.directory, tags_from_path=not args.no_tags)
        print(msg)
    elif args.cmd == "tags":
        ok, tags = engine.list_tags()
        for t in tags:
            print(f"  {t['tag']:<30} ({t['count']})")
    elif args.cmd == "dupes":
        ok, dupes = engine.find_duplicates()
        for d in dupes:
            print(f"  [{d['count']}x] {d['hash'][:12]}...")
            for p in d['paths']:
                print(f"    {p}")
            print()
    else:
        parser.print_help()

    return 0


if __name__ == "__main__":
    sys.exit(main())
